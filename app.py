import os
import sys
import json
import csv
import io
import time
import datetime
import threading
import subprocess
import urllib.request
import requests
import psutil
from typing import Optional, List, Dict, Any

import uvicorn
from fastapi import FastAPI, BackgroundTasks, HTTPException, Query, Response
from fastapi.staticfiles import StaticFiles
from fastapi.responses import HTMLResponse, StreamingResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

# Internal imports
try:
    from . import db, crawler, matcher, extractor, llm_extractor
except (ImportError, ValueError):
    import db, crawler, matcher, extractor, llm_extractor

# Safe stream fallback for windowed / frozen executables
class SafeStream:
    def write(self, s):
        pass
    def flush(self):
        pass
    def isatty(self):
        return False

if sys.stdout is None:
    sys.stdout = SafeStream()
elif hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

if sys.stderr is None:
    sys.stderr = SafeStream()
elif hasattr(sys.stderr, "reconfigure"):
    try:
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

app = FastAPI(title="FilmFreeway Intel Web UI", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

import shutil

if getattr(sys, "frozen", False):
    EXE_DIR = os.path.dirname(sys.executable)
    BUNDLE_DIR = getattr(sys, "_MEIPASS", EXE_DIR)
    STATIC_DIR = os.path.join(BUNDLE_DIR, "static")
    if not os.path.exists(STATIC_DIR):
        STATIC_DIR = os.path.join(EXE_DIR, "static")

    DB_PATH = os.path.join(EXE_DIR, "filmfreeway.db")
    if not os.path.exists(DB_PATH):
        bundle_db = os.path.join(BUNDLE_DIR, "filmfreeway.db")
        if os.path.exists(bundle_db):
            shutil.copy2(bundle_db, DB_PATH)

    PROFILE_PATH = os.path.join(EXE_DIR, "profile.json")
    if not os.path.exists(PROFILE_PATH):
        bundle_prof = os.path.join(BUNDLE_DIR, "profile.json")
        if os.path.exists(bundle_prof):
            shutil.copy2(bundle_prof, PROFILE_PATH)

    SETTINGS_PATH = os.path.join(EXE_DIR, "settings.json")
    BASE_DIR = EXE_DIR
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
    STATIC_DIR = os.path.join(BASE_DIR, "static")
    PROFILE_PATH = os.path.join(BASE_DIR, "profile.json")
    DB_PATH = os.path.join(BASE_DIR, "filmfreeway.db")
    SETTINGS_PATH = os.path.join(BASE_DIR, "settings.json")


CURATED_MODELS = [
    {
        "name": "qwen2.5:0.5b",
        "display_name": "Qwen 2.5 (0.5B) — Ultra Léger",
        "family": "Qwen",
        "parameters": "490M",
        "disk_size": "398 Mo",
        "ram_usage": "~600 Mo",
        "speed_core_i3": "35 tok/s (Instantané)",
        "accuracy_rate": "91%",
        "tier_badge": "Recommandé Core i3",
        "tier_color": "emerald",
        "description": "Modèle ultra-compact idéal pour votre Core i3 et 8 Go de RAM. Zéro ralentissement, extraction JSON immédiate.",
    },
    {
        "name": "qwen2.5:1.5b",
        "display_name": "Qwen 2.5 (1.5B) — Équilibre Idéal",
        "family": "Qwen",
        "parameters": "1.5B",
        "disk_size": "986 Mo",
        "ram_usage": "~1.8 Go",
        "speed_core_i3": "18 tok/s (Fluide)",
        "accuracy_rate": "95%",
        "tier_badge": "Équilibré & Précis",
        "tier_color": "indigo",
        "description": "Compréhension approfondie des règlements complexes. Recommandé si vous n'avez pas trop d'onglets de navigateur ouverts.",
    },
    {
        "name": "llama3.2:1b",
        "display_name": "Llama 3.2 (1B) — Meta AI Compact",
        "family": "Llama",
        "parameters": "1.2B",
        "disk_size": "1.3 Go",
        "ram_usage": "~1.5 Go",
        "speed_core_i3": "20 tok/s (Rapide)",
        "accuracy_rate": "93%",
        "tier_badge": "Léger & Robuste",
        "tier_color": "blue",
        "description": "Le petit modèle officiel de Meta. Très robuste sur les nuances de festivals anglophones.",
    },
    {
        "name": "deepseek-r1:1.5b",
        "display_name": "DeepSeek R1 (1.5B) — Raisonnement",
        "family": "DeepSeek",
        "parameters": "1.5B",
        "disk_size": "1.1 Go",
        "ram_usage": "~2.0 Go",
        "speed_core_i3": "14 tok/s (Réfléchi)",
        "accuracy_rate": "96%",
        "tier_badge": "Raisonnement Avancé",
        "tier_color": "purple",
        "description": "Modèle avec chaîne de pensée pas à pas (chain-of-thought) pour démêler les conditions d'admissibilité les plus pièges.",
    },
    {
        "name": "mistral:7b",
        "display_name": "Mistral (7B) — Expert Puissant",
        "family": "Mistral",
        "parameters": "7B",
        "disk_size": "4.4 Go",
        "ram_usage": "~5.5 Go",
        "speed_core_i3": "3 tok/s (Lourd sur CPU)",
        "accuracy_rate": "98%",
        "tier_badge": "⚠️ Attention RAM",
        "tier_color": "rose",
        "description": "Très haute précision mais volumineux. Déconseillé sur vos 8 Go de RAM totale sauf si aucun autre programme ne tourne.",
    },
]

def get_settings() -> Dict[str, Any]:
    default_settings = {
        "active_model": "qwen2.5:0.5b",
        "auto_unload_ram": True,
    }
    if os.path.exists(SETTINGS_PATH):
        try:
            with open(SETTINGS_PATH, "r", encoding="utf-8") as f:
                d = json.load(f)
                default_settings.update(d)
        except Exception:
            pass
    return default_settings

def save_settings(settings: Dict[str, Any]):
    try:
        with open(SETTINGS_PATH, "w", encoding="utf-8") as f:
            json.dump(settings, f, indent=2, ensure_ascii=False)
    except Exception:
        pass

class ModelPullManager:
    def __init__(self):
        self.lock = threading.Lock()
        self.is_pulling = False
        self.model_name: Optional[str] = None
        self.status = "idle"
        self.percent = 0
        self.completed_bytes = 0
        self.total_bytes = 0
        self.error: Optional[str] = None

    def start(self, model_name: str):
        with self.lock:
            self.is_pulling = True
            self.model_name = model_name
            self.status = f"Connexion et démarrage du téléchargement de {model_name}..."
            self.percent = 0
            self.completed_bytes = 0
            self.total_bytes = 0
            self.error = None

    def update(self, status: str, completed: int = 0, total: int = 0):
        with self.lock:
            self.status = status
            self.completed_bytes = completed
            self.total_bytes = total
            if total > 0:
                self.percent = min(100, max(1, int((completed / total) * 100)))

    def finish(self, success_msg: str):
        with self.lock:
            self.is_pulling = False
            self.percent = 100
            self.status = success_msg

    def fail(self, err_msg: str):
        with self.lock:
            self.is_pulling = False
            self.error = err_msg
            self.status = f"Erreur: {err_msg}"

    def get_status(self) -> Dict[str, Any]:
        with self.lock:
            return {
                "is_pulling": self.is_pulling,
                "model_name": self.model_name,
                "status": self.status,
                "percent": self.percent,
                "completed_bytes": self.completed_bytes,
                "total_bytes": self.total_bytes,
                "error": self.error,
            }

model_pull_mgr = ModelPullManager()

def _run_model_pull_worker(model_name: str):
    model_pull_mgr.start(model_name)
    try:
        url = "http://localhost:11434/api/pull"
        resp = requests.post(url, json={"name": model_name}, stream=True, timeout=900)
        if resp.status_code != 200:
            model_pull_mgr.fail(f"Ollama a retourné une erreur HTTP {resp.status_code}")
            return

        for line in resp.iter_lines():
            if not line:
                continue
            try:
                data = json.loads(line.decode("utf-8"))
                status = data.get("status", "")
                completed = data.get("completed", 0)
                total = data.get("total", 0)
                if status == "success":
                    model_pull_mgr.finish(f"✓ Modèle {model_name} téléchargé et prêt !")
                    return
                model_pull_mgr.update(status, completed, total)
            except Exception:
                pass
        model_pull_mgr.finish(f"✓ Modèle {model_name} installé avec succès !")
    except Exception as e:
        model_pull_mgr.fail(str(e))


# Background Task State Tracker
class TaskManager:
    def __init__(self):
        self.lock = threading.Lock()
        self.is_running = False
        self.task_type = "idle"  # 'crawl' or 'llm_review'
        self.progress = 0
        self.message = "Prêt"
        self.logs: List[str] = []
        self.error: Optional[str] = None

    def start(self, task_type: str, initial_msg: str):
        with self.lock:
            self.is_running = True
            self.task_type = task_type
            self.progress = 0
            self.message = initial_msg
            self.logs = [f"[{self._now()}] {initial_msg}"]
            self.error = None

    def log(self, text: str, progress: Optional[int] = None):
        with self.lock:
            self.logs.append(f"[{self._now()}] {text}")
            if len(self.logs) > 500:
                self.logs.pop(0)
            if progress is not None:
                self.progress = progress
            self.message = text

    def finish(self, success_msg: str):
        with self.lock:
            self.is_running = False
            self.progress = 100
            self.message = success_msg
            self.logs.append(f"[{self._now()}] ✓ {success_msg}")

    def fail(self, err_msg: str):
        with self.lock:
            self.is_running = False
            self.error = err_msg
            self.message = f"Erreur: {err_msg}"
            self.logs.append(f"[{self._now()}] ✗ Erreur: {err_msg}")

    def status(self) -> Dict[str, Any]:
        with self.lock:
            return {
                "is_running": self.is_running,
                "task_type": self.task_type,
                "progress": self.progress,
                "message": self.message,
                "logs": self.logs[-60:],  # return last 60 lines
                "error": self.error,
            }

    @staticmethod
    def _now():
        return datetime.datetime.now().strftime("%H:%M:%S")

task_mgr = TaskManager()


# Request Models (5 Comprehensive Categories)
class ProfileModel(BaseModel):
    # 1. Profil du Réalisateur & de l'Équipe
    director_age_group: str = "adult"  # youth (<18), young_adult (18-25), adult
    is_student: bool = False
    school_name: Optional[str] = None
    shot_during_studies: bool = False
    is_first_time: bool = False
    director_nationality: str = "France"
    director_residence: str = "France"
    demographics: List[str] = []  # female_filmmaker, lgbtq, indigenous, regional_minority
    credited_roles: str = "director_only"

    # 2. Fiche Technique & Artistique
    film_title: str = "Mon Court Métrage"
    runtime_minutes: int = 15
    project_type: str = "fiction"  # fiction, documentary, animation, experimental, music_video, web_series, vr_360, art_installation
    genres: List[str] = ["Short", "Drama"]
    categories: Optional[List[str]] = None
    completion_date: str = "2025-06-01"
    production_budget: Optional[float] = 5000.0
    budget_currency: str = "EUR"
    budget_tier: str = "micro"
    country_of_origin: str = "France"
    country: Optional[str] = None
    country_of_filming: str = "France"
    spoken_language: str = "Français"
    subtitles: List[str] = ["English"]
    color_type: str = "color"
    aspect_ratio: str = "16:9"
    capture_format: str = "digital"

    # 3. Politiques IA
    ai_integration_level: str = "none"  # none (0%), hybrid, full (100%)
    ai_target_category: str = "all"
    has_ai_disclosure: bool = False

    # 4. Statut de diffusion & Premières
    premiere_status_available: str = "world"  # world, international, national, regional, none
    online_availability: str = "private"  # private, public
    past_screenings_awards: Optional[str] = None
    commercial_distribution: str = "none"

    # 5. Cible & Filtres Festivals
    max_fee_usd: Optional[float] = 25.0
    deadline_not_before: Optional[str] = None
    min_years_running: int = 0
    min_reviews: int = 0
    target_accreditations: List[str] = []
    target_event_types: List[str] = ["live_screening"]
    target_region: str = "all"


class DeepSearchRequest(ProfileModel):
    search_depth: str = "standard"  # "fast", "standard", "deep"
    custom_keywords: Optional[str] = None
    delay: float = 1.5
    selected_llm_model: str = "qwen2.5:0.5b"


class CrawlRequest(BaseModel):
    query: str = "short film"
    max_pages: int = 2
    max_details: Optional[int] = 10
    delay: float = 1.5


STATUS_MAPPING = {
    "À soumettre": "To Submit",
    "Dossier en cours": "In Progress",
    "Soumis": "Submitted",
    "Sélectionné": "Selected",
    "Non retenu": "Not Selected",
    "To Submit": "To Submit",
    "In Progress": "In Progress",
    "Submitted": "Submitted",
    "Selected": "Selected",
    "Not Selected": "Not Selected",
}


class BookmarkRequest(BaseModel):
    notes: Optional[str] = ""
    submission_status: Optional[str] = "To Submit"


class BookmarkUpdateRequest(BaseModel):
    notes: Optional[str] = None
    submission_status: Optional[str] = None


class AnalysisSaveRequest(BaseModel):
    name: str
    note: Optional[str] = ""


class AnalysisImportRequest(BaseModel):
    name: Optional[str] = None
    film_title: Optional[str] = None
    profile: Dict[str, Any]
    results: List[Dict[str, Any]]
    note: Optional[str] = ""


class AISummaryRequest(BaseModel):
    model: Optional[str] = None



def get_current_profile() -> Dict[str, Any]:
    default_prof = {
        "director_age_group": "adult",
        "is_student": False,
        "school_name": "",
        "shot_during_studies": False,
        "is_first_time": False,
        "director_nationality": "France",
        "director_residence": "France",
        "demographics": [],
        "credited_roles": "director_only",
        "film_title": "Mon Court Métrage",
        "runtime_minutes": 15,
        "project_type": "fiction",
        "genres": ["Short", "Drama"],
        "categories": ["Short", "Drama"],
        "completion_date": "2025-06-01",
        "production_budget": 5000.0,
        "budget_currency": "EUR",
        "budget_tier": "micro",
        "country_of_origin": "France",
        "country": "France",
        "country_of_filming": "France",
        "spoken_language": "Français",
        "subtitles": ["English"],
        "color_type": "color",
        "aspect_ratio": "16:9",
        "capture_format": "digital",
        "ai_integration_level": "none",
        "ai_target_category": "all",
        "has_ai_disclosure": False,
        "premiere_status_available": "world",
        "online_availability": "private",
        "past_screenings_awards": "",
        "commercial_distribution": "none",
        "max_fee_usd": 25.0,
        "deadline_not_before": datetime.date.today().isoformat(),
        "min_years_running": 0,
        "min_reviews": 0,
        "target_accreditations": [],
        "target_event_types": ["live_screening"],
        "target_region": "all",
    }
    if os.path.exists(PROFILE_PATH):
        try:
            with open(PROFILE_PATH, "r", encoding="utf-8") as f:
                d = json.load(f)
                default_prof.update(d)
                # Sync categories and genres
                if "categories" in d and not d.get("genres"):
                    default_prof["genres"] = d["categories"]
                elif "genres" in d and not d.get("categories"):
                    default_prof["categories"] = d["genres"]
                return default_prof
        except Exception:
            pass
    return default_prof


import socket

def check_ollama_alive() -> Dict[str, Any]:
    try:
        # Fast socket probe in 0.15s to avoid DNS/HTTP timeout when offline
        with socket.create_connection(("127.0.0.1", 11434), timeout=0.15):
            pass
        req = urllib.request.Request("http://127.0.0.1:11434/api/version", headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=0.5) as resp:
            data = json.loads(resp.read().decode())
            return {"online": True, "version": data.get("version", "unknown")}
    except Exception:
        return {"online": False, "version": None}


# Fast Health Endpoint
@app.get("/api/health")
def api_health():
    return {"status": "ok", "app": "FilmFreeway DeepSearch Pro"}


# Clean Shutdown Endpoint
@app.post("/api/shutdown")
def api_shutdown(background_tasks: BackgroundTasks):
    def kill_proc():
        time.sleep(0.8)
        os._exit(0)
    background_tasks.add_task(kill_proc)
    return {"status": "shutting_down", "message": "Server stopping cleanly."}


# API Endpoints
@app.get("/api/stats")

def get_stats():
    conn = db.get_conn(DB_PATH)
    total = conn.execute("SELECT COUNT(*) FROM festivals").fetchone()[0]
    detailed = conn.execute("SELECT COUNT(*) FROM festivals WHERE detail_json IS NOT NULL").fetchone()[0]
    needs_review = conn.execute("SELECT COUNT(*) FROM festivals WHERE needs_review=1").fetchone()[0]
    avg_conf_row = conn.execute("SELECT AVG(confidence) FROM festivals WHERE detail_json IS NOT NULL").fetchone()[0]
    avg_conf = round(avg_conf_row or 0.0, 2)
    by_llm = conn.execute("SELECT COUNT(*) FROM festivals WHERE extracted_by='llm'").fetchone()[0]
    by_regex = conn.execute("SELECT COUNT(*) FROM festivals WHERE extracted_by='regex'").fetchone()[0]
    conn.close()

    ollama = check_ollama_alive()

    # Memory info
    free_ram_mb = None
    try:
        import psutil
        free_ram_mb = round(psutil.virtual_memory().available / (1024 * 1024), 1)
    except Exception:
        pass

    return {
        "total_festivals": total,
        "detailed_festivals": detailed,
        "needs_review": needs_review,
        "avg_confidence": avg_conf,
        "extracted_by_regex": by_regex,
        "extracted_by_llm": by_llm,
        "ollama": ollama,
        "free_ram_mb": free_ram_mb,
    }


@app.get("/api/profile")
def api_get_profile():
    return get_current_profile()


@app.post("/api/profile")
def api_save_profile(profile: ProfileModel):
    with open(PROFILE_PATH, "w", encoding="utf-8") as f:
        json.dump(profile.model_dump(), f, ensure_ascii=False, indent=2)
    return {"status": "ok", "profile": profile.model_dump()}


def get_filtered_and_sorted_festivals(
    search: Optional[str] = None,
    only_within_budget: bool = False,
    only_upcoming: bool = False,
    eligible_only: bool = False,
    min_score: Optional[int] = None,
    max_fee: Optional[float] = None,
    country: Optional[str] = None,
    live_only: bool = False,
    accreditations_only: bool = False,
    saved_only: bool = False,
    slugs: Optional[str] = None,
    sort_by: str = "score_desc"
):
    profile = get_current_profile()
    conn = db.get_conn(DB_PATH)
    rows = db.get_all_with_detail(conn)
    saved_list = db.get_saved_festivals(conn)
    conn.close()

    saved_map = {s["slug"]: s for s in saved_list}
    results = matcher.match_festivals(rows, profile)

    # Attach saved metadata to every festival
    for r in results:
        sl = r["slug"]
        if sl in saved_map:
            r["is_saved"] = True
            r["saved_notes"] = saved_map[sl].get("notes", "")
            raw_st = saved_map[sl].get("submission_status", "To Submit")
            r["submission_status"] = STATUS_MAPPING.get(raw_st, raw_st)
            r["saved_at"] = saved_map[sl].get("saved_at", "")
        else:
            r["is_saved"] = False
            r["saved_notes"] = ""
            r["submission_status"] = "To Submit"
            r["saved_at"] = None

    # Explicit slugs filter (if specified)
    if slugs:
        allowed_slugs = set(s.strip() for s in slugs.split(",") if s.strip())
        results = [r for r in results if r["slug"] in allowed_slugs]

    total_matched = len(results)

    # Dynamic multi-criteria filtering
    filtered = []
    for r in results:
        # Search query matching (name, slug, location, reasons, accreditations)
        if search:
            q = search.lower().strip()
            name = (r.get("name") or "").lower()
            slug = (r.get("slug") or "").lower()
            location = (r.get("location") or "").lower()
            why = " ".join(r.get("why_matched") or []).lower()
            accreds = " ".join(r.get("accreditations") or []).lower()
            if q not in name and q not in slug and q not in location and q not in why and q not in accreds:
                continue

        if eligible_only and not r.get("is_eligible", True):
            continue
        if only_within_budget and not r.get("fee_within_budget"):
            continue
        if only_upcoming and not r.get("has_upcoming_deadline"):
            continue
        if min_score is not None and r.get("match_pct", 0) < min_score:
            continue
        if max_fee is not None:
            fee = r.get("min_fee_usd")
            if fee is not None and fee > max_fee:
                continue
        if country and country.strip():
            c_low = country.strip().lower()
            loc_low = (r.get("location") or "").lower()
            if c_low not in loc_low:
                continue
        if live_only and not r.get("is_live_screening"):
            continue
        if accreditations_only and not r.get("accreditations"):
            continue
        if saved_only and not r.get("is_saved"):
            continue

        filtered.append(r)

    # Dynamic Sorting
    if sort_by == "score_desc":
        filtered.sort(key=lambda x: (x.get("is_eligible", True), x["score"]), reverse=True)
    elif sort_by == "score_asc":
        filtered.sort(key=lambda x: (x.get("is_eligible", True), x["score"]))
    elif sort_by == "fee_asc":
        filtered.sort(key=lambda x: (x["min_fee_usd"] is None, x["min_fee_usd"] or 999999))
    elif sort_by == "fee_desc":
        filtered.sort(key=lambda x: (x["min_fee_usd"] is not None, x["min_fee_usd"] or -1), reverse=True)
    elif sort_by == "deadline_asc":
        def dl_key(x):
            d_str = x.get("next_deadline_date")
            parsed = matcher._try_parse_date(d_str) if d_str else None
            return (parsed is None, parsed or datetime.date(2099, 1, 1))
        filtered.sort(key=dl_key)
    elif sort_by == "trust_desc":
        filtered.sort(key=lambda x: x["trust_score"], reverse=True)
    elif sort_by == "name_asc":
        filtered.sort(key=lambda x: (x["name"] or x["slug"]).lower())

    return filtered, profile, total_matched


@app.get("/api/festivals")
def api_get_festivals(
    search: Optional[str] = None,
    only_within_budget: bool = False,
    only_upcoming: bool = False,
    eligible_only: bool = False,
    min_score: Optional[int] = None,
    max_fee: Optional[float] = None,
    country: Optional[str] = None,
    live_only: bool = False,
    accreditations_only: bool = False,
    saved_only: bool = False,
    slugs: Optional[str] = None,
    sort_by: str = "score_desc"
):
    filtered, profile, total = get_filtered_and_sorted_festivals(
        search=search,
        only_within_budget=only_within_budget,
        only_upcoming=only_upcoming,
        eligible_only=eligible_only,
        min_score=min_score,
        max_fee=max_fee,
        country=country,
        live_only=live_only,
        accreditations_only=accreditations_only,
        saved_only=saved_only,
        slugs=slugs,
        sort_by=sort_by
    )

    return {
        "count": len(filtered),
        "total": total,
        "results": filtered,
        "profile": profile,
    }


@app.get("/api/festivals/{slug}")
def api_get_festival_detail(slug: str):
    conn = db.get_conn(DB_PATH)
    row = conn.execute("SELECT * FROM festivals WHERE slug=?", (slug,)).fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Festival not found")

    d = dict(row)
    detail_data = json.loads(d["detail_json"]) if d["detail_json"] else {}
    d["detail_json"] = detail_data
    bookmark_info = db.is_festival_saved(conn, slug)
    conn.close()

    profile = get_current_profile()
    matched_single = matcher.match_festivals([d], profile)
    match_info = matched_single[0] if matched_single else {}

    # 1. Évaluation détaillée catégorie par catégorie
    raw_cats = detail_data.get("categories", [])
    evaluated_cats = matcher.evaluate_all_categories(raw_cats, profile)

    # 2. Exigences et règles structurées
    constraints = detail_data.get("constraints") or {}
    requirements = {
        "max_runtime_mins": constraints.get("max_short_runtime_mins"),
        "min_completion_year": constraints.get("min_completion_year"),
        "premiere_required": constraints.get("premiere_required", "none"),
        "ai_banned": bool(constraints.get("ai_banned")),
        "has_ai_category": bool(constraints.get("has_ai_category")),
        "ai_disclosure_required": bool(constraints.get("ai_disclosure_required")),
        "public_online_banned": bool(constraints.get("public_online_banned")),
        "english_subtitles_mandatory": bool(constraints.get("english_subtitles_mandatory")),
        "is_live_screening": bool(constraints.get("is_live_screening")),
        "is_online_festival": bool(constraints.get("is_online_festival")),
        "is_contest_only": bool(constraints.get("is_contest_only")),
        "accreditations": match_info.get("accreditations", []),
    }

    # 3. Synthèse analytique & IA
    fest_info_for_summary = {
        "slug": slug,
        "name": d["name"] or detail_data.get("name") or slug,
        "detail": detail_data,
        "match": match_info,
        "location": d["listing_location"] or detail_data.get("location"),
        "years_running": d["listing_years_running"] or detail_data.get("years_running"),
        "accreditations": match_info.get("accreditations", []),
    }
    ai_summary_obj = llm_extractor.generate_ai_festival_summary(fest_info_for_summary, profile)
    ai_summary_text = ai_summary_obj.get("content", "") if isinstance(ai_summary_obj, dict) else str(ai_summary_obj)
    ai_source = ai_summary_obj.get("source", "analyse_heuristique_structuree") if isinstance(ai_summary_obj, dict) else "IA"

    return {
        "slug": d["slug"],
        "url": d["url"],
        "name": d["name"] or detail_data.get("name") or d["slug"],
        "listing_status": d["listing_status"],
        "listing_date": d["listing_date"],
        "listing_location": d["listing_location"] or detail_data.get("location"),
        "listing_years_running": d["listing_years_running"] or detail_data.get("years_running"),
        "listing_reviews": d["listing_reviews"],
        "listing_sponsored": bool(d["listing_sponsored"]),
        "confidence": d["confidence"],
        "needs_review": bool(d["needs_review"]),
        "review_reason": d["review_reason"],
        "extracted_by": d["extracted_by"],
        "last_scraped_detail": d["last_scraped_detail"],
        "detail": detail_data,
        "match": match_info,
        "evaluated_categories": evaluated_cats,
        "requirements": requirements,
        "ai_summary": ai_summary_text,
        "ai_source": ai_source,
        "ai_summary_obj": ai_summary_obj,
        "bookmark": bookmark_info,
    }


@app.post("/api/festivals/{slug}/ai-summary")
def api_generate_ai_summary(slug: str, req: Optional[AISummaryRequest] = None):
    conn = db.get_conn(DB_PATH)
    row = conn.execute("SELECT * FROM festivals WHERE slug=?", (slug,)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(status_code=404, detail="Festival introuvable")

    d = dict(row)
    detail_data = json.loads(d["detail_json"]) if d["detail_json"] else {}
    profile = get_current_profile()
    matched_single = matcher.match_festivals([d], profile)
    match_info = matched_single[0] if matched_single else {}

    settings = get_settings()
    selected_model = (req and req.model) or settings.get("active_model", "qwen2.5:0.5b")

    fest_info = {
        "slug": slug,
        "name": d["name"] or detail_data.get("name") or slug,
        "detail": detail_data,
        "match": match_info,
        "location": d["listing_location"] or detail_data.get("location"),
        "years_running": d["listing_years_running"] or detail_data.get("years_running"),
        "accreditations": match_info.get("accreditations", []),
    }
    summary_obj = llm_extractor.generate_ai_festival_summary(fest_info, profile, model=selected_model)
    ai_text = summary_obj.get("content", "") if isinstance(summary_obj, dict) else str(summary_obj)
    source = summary_obj.get("source", "analyse_heuristique_structuree") if isinstance(summary_obj, dict) else "IA"
    return {
        "status": "success",
        "summary": ai_text,
        "ai_summary": ai_text,
        "source": source,
        "raw": summary_obj,
    }



# ---------------------------------------------------------------------------
# FAVORIS / FESTIVALS SAUVEGARDÉS (WATCHLIST) ENDPOINTS
# ---------------------------------------------------------------------------

@app.get("/api/bookmarks")
def api_get_bookmarks():
    conn = db.get_conn(DB_PATH)
    saved_items = db.get_saved_festivals(conn)
    conn.close()

    profile = get_current_profile()
    if not saved_items:
        return {"count": 0, "bookmarks": [], "profile": profile}

    matched = matcher.match_festivals(saved_items, profile)
    saved_map = {s["slug"]: s for s in saved_items}
    for m in matched:
        sl = m["slug"]
        m["is_saved"] = True
        m["saved_notes"] = saved_map[sl].get("notes", "")
        raw_st = saved_map[sl].get("submission_status", "To Submit")
        m["submission_status"] = STATUS_MAPPING.get(raw_st, raw_st)
        m["saved_at"] = saved_map[sl].get("saved_at", "")

    return {"count": len(matched), "bookmarks": matched, "profile": profile}


@app.post("/api/bookmarks/{slug}")
def api_add_bookmark(slug: str, req: Optional[BookmarkRequest] = None):
    conn = db.get_conn(DB_PATH)
    row = conn.execute("SELECT slug FROM festivals WHERE slug=?", (slug,)).fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Festival not found")

    notes = req.notes if req else ""
    raw_status = req.submission_status if req else "To Submit"
    status = STATUS_MAPPING.get(raw_status, raw_status)
    db.save_festival_bookmark(conn, slug, notes=notes, submission_status=status)
    conn.close()
    return {"status": "success", "slug": slug, "saved": True}


@app.delete("/api/bookmarks/{slug}")
def api_remove_bookmark(slug: str):
    conn = db.get_conn(DB_PATH)
    db.remove_festival_bookmark(conn, slug)
    conn.close()
    return {"status": "success", "slug": slug, "saved": False}


@app.patch("/api/bookmarks/{slug}")
def api_update_bookmark(slug: str, req: BookmarkUpdateRequest):
    conn = db.get_conn(DB_PATH)
    db.update_saved_festival(conn, slug, notes=req.notes, submission_status=req.submission_status)
    info = db.is_festival_saved(conn, slug)
    conn.close()
    return {"status": "success", "slug": slug, **info}


# ---------------------------------------------------------------------------
# HISTORIQUE DES ANALYSES & SESSIONS ENDPOINTS
# ---------------------------------------------------------------------------

@app.get("/api/analyses")
def api_list_analyses():
    conn = db.get_conn(DB_PATH)
    snapshots = db.list_analysis_snapshots(conn)
    conn.close()
    return {"count": len(snapshots), "analyses": snapshots}


@app.post("/api/analyses/save")
def api_save_analysis(req: AnalysisSaveRequest):
    profile = get_current_profile()
    conn = db.get_conn(DB_PATH)
    rows = db.get_all_with_detail(conn)
    results = matcher.match_festivals(rows, profile)

    film_title = profile.get("film_title", "Court-Métrage")
    name = req.name.strip() if req.name and req.name.strip() else f"Analyse {film_title} ({datetime.date.today().strftime('%d/%m/%Y')})"

    snap_id = db.save_analysis_snapshot(
        conn,
        name=name,
        film_title=film_title,
        profile=profile,
        results=results,
        note=req.note or ""
    )
    conn.close()
    return {"status": "success", "id": snap_id, "name": name}


@app.get("/api/analyses/{snapshot_id}")
def api_get_analysis(snapshot_id: int):
    conn = db.get_conn(DB_PATH)
    snapshot = db.get_analysis_snapshot(conn, snapshot_id)
    saved_list = db.get_saved_festivals(conn)
    conn.close()
    if not snapshot:
        raise HTTPException(status_code=404, detail="Analyse introuvable")

    saved_map = {s["slug"]: s for s in saved_list}
    for r in snapshot.get("results", []):
        sl = r.get("slug")
        if sl in saved_map:
            r["is_saved"] = True
            r["saved_notes"] = saved_map[sl].get("notes", "")
            raw_st = saved_map[sl].get("submission_status", "To Submit")
            r["submission_status"] = STATUS_MAPPING.get(raw_st, raw_st)
        else:
            r["is_saved"] = False
            r["saved_notes"] = ""
            r["submission_status"] = "To Submit"

    return snapshot


@app.delete("/api/analyses/{snapshot_id}")
def api_delete_analysis(snapshot_id: int):
    conn = db.get_conn(DB_PATH)
    db.delete_analysis_snapshot(conn, snapshot_id)
    conn.close()
    return {"status": "success", "id": snapshot_id}


@app.get("/api/analyses/{snapshot_id}/export-file")
def api_export_analysis_file(snapshot_id: int):
    conn = db.get_conn(DB_PATH)
    snapshot = db.get_analysis_snapshot(conn, snapshot_id)
    conn.close()
    if not snapshot:
        raise HTTPException(status_code=404, detail="Analyse introuvable")

    filename = f"analyse_{snapshot_id}_{snapshot.get('film_title', 'court').replace(' ', '_')}.json"
    return Response(
        content=json.dumps(snapshot, ensure_ascii=False, indent=2),
        media_type="application/json; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@app.post("/api/analyses/import-session")
def api_import_session(data: AnalysisImportRequest):
    conn = db.get_conn(DB_PATH)
    name = data.name or f"Import {data.film_title or 'Court'} ({datetime.date.today().strftime('%d/%m/%Y')})"
    snap_id = db.save_analysis_snapshot(
        conn,
        name=name,
        film_title=data.film_title or data.profile.get("film_title", "Court"),
        profile=data.profile,
        results=data.results,
        note=data.note or "Importé depuis un fichier session"
    )
    conn.close()
    return {"status": "success", "id": snap_id, "name": name}



# Background Deep Search Worker
def _run_deep_search_worker(
    prof_dict: Dict[str, Any],
    search_depth: str,
    custom_keywords: Optional[str],
    delay: float,
    selected_llm_model: str = "qwen2.5:0.5b"
):
    try:
        # Save updated film profile
        with open(PROFILE_PATH, "w", encoding="utf-8") as f:
            json.dump(prof_dict, f, ensure_ascii=False, indent=2)

        film_title = prof_dict.get("film_title", "Mon Court Métrage")
        task_mgr.start("deep_search", f"Initialisation de la recherche approfondie pour '{film_title}'...")

        # Build smart queries targeting short film competitions for these categories
        queries = []
        if custom_keywords and custom_keywords.strip():
            queries.append(custom_keywords.strip())

        # Project type query
        p_type = prof_dict.get("project_type", "fiction").lower()
        if p_type == "documentary":
            queries.append("short documentary")
        elif p_type == "animation":
            queries.append("short animation")
        elif p_type == "experimental":
            queries.append("experimental short")
        elif p_type == "music_video":
            queries.append("music video")

        # Genres
        genres = prof_dict.get("genres") or prof_dict.get("categories") or ["Short"]
        for g in genres:
            g_clean = g.strip().lower()
            if g_clean in ("short", "court métrage", "court metrage"):
                if "short film" not in queries:
                    queries.append("short film")
            else:
                q_candidate = f"short {g_clean}"
                if q_candidate not in queries:
                    queries.append(q_candidate)

        if prof_dict.get("is_student") and "student film" not in queries:
            queries.append("student film")

        if prof_dict.get("ai_integration_level") == "full" and "ai film" not in queries:
            queries.append("ai film")

        if not queries:
            queries = ["short film"]

        # Limit to 3 queries max to stay polite and fast
        queries = queries[:3]

        if search_depth == "fast":
            pages_per_query = 2
            max_details = 15
        elif search_depth == "deep":
            pages_per_query = 5
            max_details = 60
        else:  # "standard"
            pages_per_query = 3
            max_details = 30

        conn = db.get_conn(DB_PATH)
        session = crawler.create_session()
        total_found = 0
        newly_found_slugs = []

        # Step 1: Listing exploration
        for q_idx, query in enumerate(queries, 1):
            task_mgr.log(
                f"[1/3] Recherche FilmFreeway : '{query}' ({q_idx}/{len(queries)})...",
                progress=int(5 + (q_idx / len(queries)) * 25)
            )
            for page in range(1, pages_per_query + 1):
                params = {"page": page, "q": query}
                resp = crawler.polite_get(session, crawler.SEARCH_URL, params=params, base_delay=delay)
                if resp is None:
                    break
                entries = extractor.parse_search_page(resp.text)
                if not entries:
                    break
                for e in entries:
                    db.upsert_listing_entry(conn, e)
                    newly_found_slugs.append(e["slug"])
                total_found += len(entries)
                task_mgr.log(f"  -> '{query}' page {page} : +{len(entries)} festivals trouvés ({total_found} cumulés).")
                tot_p = extractor.get_total_pages(resp.text)
                if tot_p and page >= tot_p:
                    break

        task_mgr.log(f"[1/3] Exploration catalogue terminée ({total_found} festivals répertoriés).", progress=32)

        # Step 2: Detail analysis (prioritizing newly found slugs)
        rows = db.get_slugs_needing_detail(conn, priority_slugs=newly_found_slugs, limit=max_details)
        total_details = len(rows)
        task_mgr.log(f"[2/3] Analyse approfondie de {total_details} fiches festivals (tarifs, durées, dates)...", progress=35)

        for idx, row in enumerate(rows, 1):
            slug, url = row["slug"], row["url"]
            pct = 35 + int((idx / max(1, total_details)) * 55)
            task_mgr.log(f"  [{idx}/{total_details}] Analyse : {slug}", progress=pct)

            resp = crawler.polite_get(session, url, base_delay=delay)
            if resp is None:
                db.mark_detail_fetch_failed(conn, slug)
                continue

            detail, confidence, needs_rev, reason, raw_text = extractor.parse_festival_detail(resp.text, url, slug)
            if needs_rev and check_ollama_alive()["online"]:
                try:
                    task_mgr.log(f"    [IA : {selected_llm_model}] Révision assistée de {slug}...")
                    llm_res = llm_extractor.extract_with_llm(raw_text, model=selected_llm_model)
                    if llm_res:
                        detail.update(llm_res)
                        needs_rev = False
                        confidence = 0.85
                        task_mgr.log(f"    [IA : ✓] {slug} résolu avec succès par {selected_llm_model}")
                except Exception:
                    pass
            db.save_detail(conn, slug, detail, confidence, needs_rev, reason, 14, raw_text)

        # Step 3: Match calculation
        task_mgr.log("[3/3] Calcul des correspondances et adéquation avec votre court-métrage...", progress=95)
        all_rows = db.get_all_with_detail(conn)
        matched = matcher.match_festivals(all_rows, prof_dict)
        conn.close()

        top_matches = sum(1 for m in matched if m.get("match_pct", 0) >= 70 and m.get("is_eligible", True))
        task_mgr.finish(
            f"Recherche terminée ! {len(matched)} festivals analysés, dont {top_matches} correspondant très fortement à votre film."
        )
    except Exception as exc:
        task_mgr.fail(f"Exception durant la recherche approfondie: {str(exc)}")


@app.post("/api/deep-search")
def api_deep_search(req: DeepSearchRequest):
    if task_mgr.is_running:
        raise HTTPException(status_code=400, detail="Une recherche est déjà en cours d'exécution.")
    
    prof_dict = req.model_dump()
    search_depth = prof_dict.pop("search_depth", "standard")
    custom_keywords = prof_dict.pop("custom_keywords", None)
    delay = prof_dict.pop("delay", 1.5)
    selected_llm_model = prof_dict.pop("selected_llm_model", "qwen2.5:0.5b")

    t = threading.Thread(
        target=_run_deep_search_worker,
        args=(prof_dict, search_depth, custom_keywords, delay, selected_llm_model),
        daemon=True
    )
    t.start()
    return {"status": "started", "task": "deep_search"}


# Background Crawl Worker
def _run_crawl_worker(query: str, max_pages: int, max_details: Optional[int], delay: float):
    try:
        task_mgr.start("crawl", f"Démarrage de la recherche : '{query}' ({max_pages} pages max)...")
        conn = db.get_conn(DB_PATH)

        # 1. Listing Crawl
        session = crawler.create_session()
        total_pages = None
        found = 0

        for page in range(1, max_pages + 1):
            params = {"page": page}
            if query:
                params["q"] = query

            task_mgr.log(f"[Listing] Scrape page {page} ({query})...", progress=int((page / (max_pages + 1)) * 30))
            resp = crawler.polite_get(session, crawler.SEARCH_URL, params=params, base_delay=delay)
            if resp is None:
                task_mgr.log(f"[Listing] Échec sur la page {page}, arrêt.")
                break

            entries = extractor.parse_search_page(resp.text)
            if not entries:
                task_mgr.log(f"[Listing] Page {page} sans festivals, fin du listing.")
                break

            for e in entries:
                db.upsert_listing_entry(conn, e)
            found += len(entries)
            task_mgr.log(f"[Listing] Page {page} : +{len(entries)} festivals trouvés ({found} total).")

            if total_pages is None:
                total_pages = extractor.get_total_pages(resp.text)
            if total_pages and page >= total_pages:
                break

        task_mgr.log(f"[Listing] Indexation terminée. {found} festivals enregistrés en base.", progress=35)

        # 2. Detail Crawl
        rows = db.get_slugs_needing_detail(conn, limit=max_details)
        total_details = len(rows)
        task_mgr.log(f"[Détail] {total_details} fiches festivals à enrichir...")

        for idx, row in enumerate(rows, 1):
            slug, url = row["slug"], row["url"]
            pct = 35 + int((idx / max(1, total_details)) * 60)
            task_mgr.log(f"[Détail {idx}/{total_details}] {slug}", progress=pct)

            resp = crawler.polite_get(session, url, base_delay=delay)
            if resp is None:
                db.mark_detail_fetch_failed(conn, slug)
                task_mgr.log(f"  [!] Échec réseau pour {slug}")
                continue

            detail, confidence, needs_rev, reason, raw_text = extractor.parse_festival_detail(resp.text, url, slug)
            db.save_detail(conn, slug, detail, confidence, needs_rev, reason, 14, raw_text)

            flag = f" [!] À vérifier: {reason}" if needs_rev else f" [✓] Confiance: {confidence}"
            task_mgr.log(f"  -> {slug}: {len(detail.get('categories', []))} cat., {len(detail.get('dates', []))} dates{flag}")

        conn.close()
        task_mgr.finish(f"Recherche terminée ! {found} festivals indexés, {total_details} fiches détaillées.")
    except Exception as exc:
        task_mgr.fail(f"Exception pendant le crawl: {str(exc)}")


@app.post("/api/crawl")
def api_start_crawl(req: CrawlRequest):
    if task_mgr.is_running:
        raise HTTPException(status_code=400, detail="Une tâche est déjà en cours d'exécution.")
    t = threading.Thread(target=_run_crawl_worker, args=(req.query, req.max_pages, req.max_details, req.delay), daemon=True)
    t.start()
    return {"status": "started", "task": "crawl"}


# Background LLM Worker
def _run_llm_worker():
    try:
        task_mgr.start("llm_review", "Démarrage de la revue IA locale (Ollama)...")
        conn = db.get_conn(DB_PATH)
        queue = db.get_review_queue_with_raw(conn)
        total = len(queue)
        task_mgr.log(f"[LLM] {total} festival(s) en attente de révision.")

        if total == 0:
            conn.close()
            task_mgr.finish("Aucun festival en attente de révision !")
            return

        fixed, still_stuck = 0, 0
        for i, item in enumerate(queue, 1):
            pct = int((i / total) * 90)
            name = item["name"] or item["slug"]
            task_mgr.log(f"[LLM {i}/{total}] Analyse de '{name}'...", progress=pct)

            res = llm_extractor.extract_with_llm(item["raw_text"])
            if res:
                db.save_llm_correction(conn, item["slug"], res)
                fixed += 1
                task_mgr.log(f"  [✓] '{name}' corrigé avec succès par l'IA !")
            else:
                still_stuck += 1
                task_mgr.log(f"  [?] Échec d'extraction IA pour '{name}'.")

        conn.close()
        task_mgr.finish(f"Revue IA terminée : {fixed} corrigé(s), {still_stuck} restant(s).")
    except Exception as exc:
        task_mgr.fail(f"Exception revue LLM: {str(exc)}")


@app.post("/api/llm-review")
def api_start_llm_review():
    if task_mgr.is_running:
        raise HTTPException(status_code=400, detail="Une tâche est déjà en cours d'exécution.")
    t = threading.Thread(target=_run_llm_worker, daemon=True)
    t.start()
    return {"status": "started", "task": "llm_review"}


@app.get("/api/task/status")
def api_task_status():
    return task_mgr.status()


@app.post("/api/ollama/unload")
@app.post("/api/models/unload")
def api_unload_ollama():
    try:
        # Stop all running models in Ollama ps
        try:
            resp = requests.get("http://localhost:11434/api/ps", timeout=3)
            if resp.status_code == 200:
                for m in resp.json().get("models", []):
                    m_name = m.get("name")
                    if m_name:
                        requests.post("http://localhost:11434/api/generate", json={"model": m_name, "keep_alive": 0}, timeout=3)
        except Exception:
            pass

        cmd = [os.path.expandvars(r"%LOCALAPPDATA%\Programs\Ollama\ollama.exe"), "stop", "qwen2.5:0.5b"]
        subprocess.run(cmd, capture_output=True, text=True, timeout=5)
        return {"status": "success", "message": "Modèles déchargés de la RAM avec succès !"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


class ModelActionRequest(BaseModel):
    name: str


@app.get("/api/models")
def api_get_models():
    settings = get_settings()
    active_model = settings.get("active_model", "qwen2.5:0.5b")

    try:
        mem = psutil.virtual_memory()
        ram_info = {
            "total_gb": round(mem.total / (1024**3), 2),
            "available_gb": round(mem.available / (1024**3), 2),
            "used_percent": mem.percent,
            "cpu_info": "Intel Core i3-1005G1 (4 vCPU)",
        }
    except Exception:
        ram_info = {"total_gb": 8.0, "available_gb": 2.0, "used_percent": 75, "cpu_info": "Intel Core i3"}

    ollama_status = check_ollama_alive()
    installed_models_map = {}
    running_models_set = set()

    if ollama_status["online"]:
        try:
            req_tags = urllib.request.Request("http://localhost:11434/api/tags")
            with urllib.request.urlopen(req_tags, timeout=2.0) as resp:
                tags_data = json.loads(resp.read().decode())
                for m in tags_data.get("models", []):
                    name = m.get("name", "")
                    clean_name = name.split(":")[0] if ":" in name and name.endswith(":latest") else name
                    installed_models_map[name] = m
                    installed_models_map[clean_name] = m
        except Exception:
            pass

        try:
            req_ps = urllib.request.Request("http://localhost:11434/api/ps")
            with urllib.request.urlopen(req_ps, timeout=2.0) as resp:
                ps_data = json.loads(resp.read().decode())
                for m in ps_data.get("models", []):
                    running_models_set.add(m.get("name", ""))
        except Exception:
            pass

    models_out = []
    seen_names = set()

    for cm in CURATED_MODELS:
        name = cm["name"]
        seen_names.add(name)
        is_inst = (name in installed_models_map)
        is_run = (name in running_models_set)
        actual_size = None
        if is_inst:
            size_bytes = installed_models_map[name].get("size", 0)
            actual_size = f"{round(size_bytes / (1024**2), 1)} Mo" if size_bytes < 1024**3 else f"{round(size_bytes / (1024**3), 2)} Go"

        models_out.append({
            **cm,
            "is_installed": is_inst,
            "is_running": is_run,
            "is_active_default": (name == active_model),
            "actual_disk_size": actual_size or cm["disk_size"],
        })

    for inst_name, inst_data in installed_models_map.items():
        if inst_name not in seen_names and not inst_name.endswith(":latest"):
            size_b = inst_data.get("size", 0)
            size_str = f"{round(size_b / (1024**2), 1)} Mo" if size_b < 1024**3 else f"{round(size_b / (1024**3), 2)} Go"
            models_out.append({
                "name": inst_name,
                "display_name": inst_name,
                "family": inst_data.get("details", {}).get("family", "Custom"),
                "parameters": inst_data.get("details", {}).get("parameter_size", "Inconnu"),
                "disk_size": size_str,
                "actual_disk_size": size_str,
                "ram_usage": "Variable",
                "speed_core_i3": "Selon modèle",
                "accuracy_rate": "Variable",
                "tier_badge": "Modèle Personnalisé",
                "tier_color": "slate",
                "description": "Modèle tiers installé.",
                "is_installed": True,
                "is_running": (inst_name in running_models_set),
                "is_active_default": (inst_name == active_model),
            })
            seen_names.add(inst_name)

    return {
        "ollama_online": ollama_status["online"],
        "ollama_version": ollama_status["version"],
        "active_model": active_model,
        "pull_status": model_pull_mgr.get_status(),
        "system_ram": ram_info,
        "models": models_out,
    }


@app.post("/api/models/pull")
def api_pull_model(req: ModelActionRequest):
    name = req.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Nom du modèle requis.")
    if model_pull_mgr.is_pulling:
        raise HTTPException(status_code=400, detail="Un téléchargement est déjà en cours.")
    t = threading.Thread(target=_run_model_pull_worker, args=(name,), daemon=True)
    t.start()
    return {"status": "started", "model": name}


@app.get("/api/models/pull-status")
def api_pull_status():
    return model_pull_mgr.get_status()


@app.delete("/api/models/{model_name}")
def api_delete_model(model_name: str):
    try:
        url = "http://localhost:11434/api/delete"
        resp = requests.delete(url, json={"name": model_name}, timeout=15)
        if resp.status_code == 200:
            return {"status": "success", "message": f"Modèle '{model_name}' supprimé avec succès !"}
        else:
            return {"status": "error", "message": f"Erreur Ollama HTTP {resp.status_code}"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


@app.post("/api/models/select")
def api_select_model(req: ModelActionRequest):
    settings = get_settings()
    settings["active_model"] = req.name.strip()
    save_settings(settings)
    return {"status": "success", "active_model": settings["active_model"]}


@app.post("/api/models/start-daemon")
def api_start_ollama_daemon():
    try:
        ollama_exe = os.path.expandvars(r"%LOCALAPPDATA%\Programs\Ollama\ollama.exe")
        subprocess.Popen([ollama_exe, "serve"], creationflags=subprocess.CREATE_NEW_PROCESS_GROUP if sys.platform == "win32" else 0)
        time.sleep(1.5)
        return {"status": "success", "message": "Serveur Ollama démarré avec succès !"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


@app.get("/api/export/excel")
def api_export_excel(
    search: Optional[str] = None,
    only_within_budget: bool = False,
    only_upcoming: bool = False,
    eligible_only: bool = False,
    min_score: Optional[int] = None,
    max_fee: Optional[float] = None,
    country: Optional[str] = None,
    live_only: bool = False,
    accreditations_only: bool = False,
    saved_only: bool = False,
    slugs: Optional[str] = None,
    sort_by: str = "score_desc"
):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    results, profile, total_count = get_filtered_and_sorted_festivals(
        search=search,
        only_within_budget=only_within_budget,
        only_upcoming=only_upcoming,
        eligible_only=eligible_only,
        min_score=min_score,
        max_fee=max_fee,
        country=country,
        live_only=live_only,
        accreditations_only=accreditations_only,
        saved_only=saved_only,
        slugs=slugs,
        sort_by=sort_by
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Watchlist" if saved_only else "Recommended Festivals"
    ws.views.sheetView[0].showGridLines = True

    # Title Banner
    ws.merge_cells("A1:P1")
    title_cell = ws["A1"]
    suffix = " — WATCHLIST FAVORITES" if saved_only else " — FILTERED INTELLIGENCE REPORT"
    title_cell.value = f"FILMFREEWAY DEEPSEARCH PRO{suffix} FOR \"{profile.get('film_title', 'Short Film').upper()}\" ({datetime.date.today().strftime('%Y-%m-%d')})"
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Headers
    headers = [
        "Rank",
        "Match Score (%)",
        "Eligibility",
        "Fit Rating",
        "Festival Name",
        "City / Country",
        "Min Fee ($)",
        "Within Budget?",
        "Open Deadline?",
        "Next Deadline",
        "Trust Score",
        "Accreditations (Oscars/BAFTA)",
        "Watchlist Status",
        "Confirmed Strengths",
        "Blockers / Disqualifications",
        "FilmFreeway URL"
    ]
    ws.append(headers)
    ws.row_dimensions[2].height = 24

    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4338CA", end_color="4338CA", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Data Rows
    for i, r in enumerate(results, 1):
        row_idx = i + 2
        fee_str = f"${r['min_fee_usd']:.2f}" if r.get("min_fee_usd") is not None else "Free / Unknown"
        fee_ok_str = "YES" if r.get("fee_within_budget") else "NO (Over budget)"
        deadline_str = "YES" if r.get("has_upcoming_deadline") else "NO / Closed"
        why_str = " • ".join(r.get("why_matched") or [])
        blockers_str = " • ".join(r.get("blockers") or [])
        accred_str = ", ".join(r.get("accreditations") or []) or "Standard"
        next_dl = r.get("next_deadline_date") or r.get("next_deadline_label") or "To verify"
        elig_str = "ELIGIBLE" if r.get("is_eligible", True) else "INELIGIBLE (Disqualified)"
        raw_fav = r.get('submission_status', 'To Submit')
        fav_str = f"⭐ {STATUS_MAPPING.get(raw_fav, raw_fav)}" if r.get("is_saved") else "Standard"

        ws.append([
            i,
            f"{r.get('match_pct', 0)}%",
            elig_str,
            r.get("match_quality", ""),
            r.get("name") or r.get("slug"),
            r.get("location") or "",
            fee_str,
            fee_ok_str,
            deadline_str,
            next_dl,
            f"{int((r.get('trust_score') or 0)*100)}%",
            accred_str,
            fav_str,
            why_str,
            blockers_str,
            r.get("url")
        ])
        ws.row_dimensions[row_idx].height = 20

        # Styles
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center")
        score_cell = ws.cell(row=row_idx, column=2)
        score_cell.alignment = Alignment(horizontal="center")
        score_cell.font = Font(name="Calibri", size=10, bold=True)
        if r.get("is_eligible", True):
            if r.get("match_pct", 0) >= 80:
                score_cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                score_cell.font = Font(name="Calibri", size=10, bold=True, color="166534")
            elif r.get("match_pct", 0) >= 60:
                score_cell.fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
                score_cell.font = Font(name="Calibri", size=10, bold=True, color="3730A3")
        else:
            score_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            score_cell.font = Font(name="Calibri", size=10, bold=True, color="991B1B")

        # Url hyperlink
        url_cell = ws.cell(row=row_idx, column=16)
        url_cell.hyperlink = r.get("url")
        url_cell.font = Font(name="Calibri", size=9, color="2563EB", underline="single")

        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col_idx).border = thin_border

    # Auto column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if len(val) > max_len and cell.row > 1:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 45)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"rapport_festivals_filtres_{datetime.date.today().isoformat()}.xlsx"
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@app.get("/api/export/csv")
def api_export_csv(
    search: Optional[str] = None,
    only_within_budget: bool = False,
    only_upcoming: bool = False,
    eligible_only: bool = False,
    min_score: Optional[int] = None,
    max_fee: Optional[float] = None,
    country: Optional[str] = None,
    live_only: bool = False,
    accreditations_only: bool = False,
    saved_only: bool = False,
    slugs: Optional[str] = None,
    sort_by: str = "score_desc"
):
    results, profile, total_count = get_filtered_and_sorted_festivals(
        search=search,
        only_within_budget=only_within_budget,
        only_upcoming=only_upcoming,
        eligible_only=eligible_only,
        min_score=min_score,
        max_fee=max_fee,
        country=country,
        live_only=live_only,
        accreditations_only=accreditations_only,
        saved_only=saved_only,
        slugs=slugs,
        sort_by=sort_by
    )

    output = io.StringIO()
    output.write("\ufeff")
    writer = csv.writer(output, delimiter=";")
    writer.writerow([
        "Rank",
        "Match Score (%)",
        "Eligibility",
        "Fit Rating",
        "Festival Name",
        "City / Country",
        "FilmFreeway URL",
        "Min Fee (USD)",
        "Within Budget?",
        "Upcoming Deadline?",
        "Next Deadline",
        "Trust Score",
        "Accreditations",
        "Watchlist Status",
        "Director Notes",
        "Confirmed Strengths",
        "Blockers / Disqualifications"
    ])

    for i, r in enumerate(results, 1):
        raw_fav = r.get("submission_status", "To Submit")
        fav_label = STATUS_MAPPING.get(raw_fav, raw_fav)
        writer.writerow([
            i,
            f"{r.get('match_pct', 0)}%",
            "ELIGIBLE" if r.get("is_eligible", True) else "INELIGIBLE",
            r.get("match_quality", ""),
            r.get("name") or r.get("slug"),
            r.get("location") or "",
            r.get("url"),
            f"{r['min_fee_usd']:.2f}" if r.get("min_fee_usd") is not None else "Unknown / Free",
            "YES" if r.get("fee_within_budget") else "NO (Over budget)",
            "YES" if r.get("has_upcoming_deadline") else "NO",
            r.get("next_deadline_date") or r.get("next_deadline_label") or "To verify",
            f"{int((r.get('trust_score') or 0)*100)}%",
            ", ".join(r.get("accreditations") or []) or "Standard",
            f"⭐ {fav_label}" if r.get("is_saved") else "No",
            r.get("saved_notes", "") or "",
            " • ".join(r.get("why_matched") or []),
            " • ".join(r.get("blockers") or [])
        ])

    filename = f"filmfreeway_festivals_export_{datetime.date.today().isoformat()}.csv"
    return Response(
        content=output.getvalue().encode("utf-8-sig"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@app.get("/api/export/json")
def api_export_json(
    search: Optional[str] = None,
    only_within_budget: bool = False,
    only_upcoming: bool = False,
    eligible_only: bool = False,
    min_score: Optional[int] = None,
    max_fee: Optional[float] = None,
    country: Optional[str] = None,
    live_only: bool = False,
    accreditations_only: bool = False,
    saved_only: bool = False,
    slugs: Optional[str] = None,
    sort_by: str = "score_desc"
):
    results, profile, total_count = get_filtered_and_sorted_festivals(
        search=search,
        only_within_budget=only_within_budget,
        only_upcoming=only_upcoming,
        eligible_only=eligible_only,
        min_score=min_score,
        max_fee=max_fee,
        country=country,
        live_only=live_only,
        accreditations_only=accreditations_only,
        saved_only=saved_only,
        slugs=slugs,
        sort_by=sort_by
    )

    data = {
        "export_date": datetime.datetime.now().isoformat(),
        "film_profile": profile,
        "filters_applied": {
            "search": search,
            "only_within_budget": only_within_budget,
            "only_upcoming": only_upcoming,
            "eligible_only": eligible_only,
            "min_score": min_score,
            "max_fee": max_fee,
            "country": country,
            "live_only": live_only,
            "accreditations_only": accreditations_only,
            "saved_only": saved_only,
            "sort_by": sort_by
        },
        "count": len(results),
        "festivals_count": len(results),
        "total_in_catalogue": total_count,
        "results": results,
    }
    filename = f"rapport_festivals_filtres_{datetime.date.today().isoformat()}.json"
    return Response(
        content=json.dumps(data, ensure_ascii=False, indent=2),
        media_type="application/json; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )



# Static files mount
if os.path.exists(STATIC_DIR):
    app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

@app.get("/")
def serve_index():
    index_file = os.path.join(STATIC_DIR, "index.html")
    if os.path.exists(index_file):
        with open(index_file, "r", encoding="utf-8") as f:
            return HTMLResponse(content=f.read())
    return HTMLResponse("<h1>FilmFreeway Intel API is running. index.html not found.</h1>")


if __name__ == "__main__":
    import webbrowser
    print("==================================================")
    print("  FilmFreeway Intel — Serveur Web Moderne")
    print("  Accessible sur : http://127.0.0.1:8000")
    print("==================================================")
    if "--no-browser" not in sys.argv:
        try:
            threading.Timer(1.2, lambda: webbrowser.open("http://127.0.0.1:8000")).start()
        except Exception:
            pass
    uvicorn.run("app:app", host="127.0.0.1", port=8000, reload=False, log_level="info")
